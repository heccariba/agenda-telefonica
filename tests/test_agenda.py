"""Pruebas de la agenda telefónica - rama con-excel.

La lógica es la misma que en la versión 1.
El guardado ahora usa data/contactos.xlsx con openpyxl:
fila 1 = cabecera, desde fila 2 = datos.
"""

from openpyxl import Workbook
from openpyxl import load_workbook

from src import agenda


def test_mostrar_menu_existe():
    assert callable(agenda.mostrar_menu)


def test_opcion_salir_es_cero():
    assert agenda.OPCION_SALIR == "0"


def test_nombre_no_vacio_es_valido():
    assert agenda.nombre_es_valido("Ana") is True


def test_nombre_vacio_no_es_valido():
    assert agenda.nombre_es_valido("   ") is False


def test_telefono_9_cifras_es_valido():
    assert agenda.telefono_es_valido("600123123") is True


def test_telefono_con_letras_no_es_valido():
    assert agenda.telefono_es_valido("600abc123") is False


def test_telefono_corto_no_es_valido():
    assert agenda.telefono_es_valido("123") is False


def test_telefono_largo_no_es_valido():
    assert agenda.telefono_es_valido("1" * 16) is False


def test_anadir_contacto_valido():
    mi_agenda = {}
    assert agenda.anadir_contacto(mi_agenda, "Ana", "600123123") is True
    assert mi_agenda == {"Ana": "600123123"}


def test_anadir_nombre_duplicado_falla():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.anadir_contacto(mi_agenda, "ana", "600999999") is False
    assert mi_agenda == {"Ana": "600123123"}


def test_anadir_numero_duplicado_falla():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.anadir_contacto(mi_agenda, "Luis", "600123123") is False
    assert "Luis" not in mi_agenda


def test_buscar_por_nombre_ignora_mayusculas():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.buscar_por_nombre(mi_agenda, "ANA") == "600123123"


def test_buscar_por_nombre_inexistente_da_none():
    assert agenda.buscar_por_nombre({}, "Ana") is None


def test_buscar_por_numero_encuentra_nombre():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.buscar_por_numero(mi_agenda, "600123123") == "Ana"


def test_buscar_por_numero_inexistente_da_none():
    assert agenda.buscar_por_numero({}, "600123123") is None


def test_eliminar_por_nombre_borra():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.eliminar_por_nombre(mi_agenda, "ana") is True
    assert mi_agenda == {}


def test_eliminar_por_nombre_inexistente_falla():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.eliminar_por_nombre(mi_agenda, "Luis") is False
    assert mi_agenda == {"Ana": "600123123"}


def test_eliminar_por_numero_borra():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.eliminar_por_numero(mi_agenda, "600123123") is True
    assert mi_agenda == {}


def test_eliminar_por_numero_inexistente_falla():
    mi_agenda = {"Ana": "600123123"}
    assert agenda.eliminar_por_numero(mi_agenda, "600000000") is False
    assert mi_agenda == {"Ana": "600123123"}


def test_contactos_ordenados_alfabetico():
    mi_agenda = {"Luis": "600111111", "Ana": "600123123"}
    nombres = [nombre for nombre, _ in agenda.contactos_ordenados(mi_agenda)]
    assert nombres == ["Ana", "Luis"]


def test_contactos_ordenados_inverso():
    mi_agenda = {"Luis": "600111111", "Ana": "600123123"}
    nombres = [n for n, _ in agenda.contactos_ordenados(mi_agenda, inverso=True)]
    assert nombres == ["Luis", "Ana"]


def test_guardar_y_cargar(tmp_path):
    mi_agenda = {"Ana": "600123123", "Luis": "600111111"}
    ruta = str(tmp_path / "contactos.xlsx")
    agenda.guardar_contactos(mi_agenda, ruta)
    assert agenda.cargar_contactos(ruta) == mi_agenda


def test_cargar_fichero_inexistente_da_vacia(tmp_path):
    ruta = str(tmp_path / "no_existe.xlsx")
    assert agenda.cargar_contactos(ruta) == {}


def test_cargar_ignora_filas_malas(tmp_path):
    ruta = str(tmp_path / "contactos.xlsx")
    libro = Workbook()
    hoja = libro.active
    hoja.append(["nombre", "telefono"])
    hoja.append(["Ana", "600123123"])
    hoja.append(["SoloNombre"])
    hoja.append(["Luis", "600111111"])
    hoja.append([None, None])
    libro.save(ruta)
    libro.close()
    assert agenda.cargar_contactos(ruta) == {"Ana": "600123123", "Luis": "600111111"}


def test_primera_fila_guarda_cabecera(tmp_path):
    ruta = str(tmp_path / "contactos.xlsx")
    agenda.guardar_contactos({"Ana": "600123123"}, ruta)
    libro = load_workbook(ruta)
    hoja = libro.active
    cabecera = [hoja.cell(row=1, column=1).value, hoja.cell(row=1, column=2).value]
    libro.close()
    assert cabecera == ["nombre", "telefono"]


def test_telefono_se_guarda_como_texto(tmp_path):
    ruta = str(tmp_path / "contactos.xlsx")
    agenda.guardar_contactos({"Ana": "600123123"}, ruta)
    libro = load_workbook(ruta)
    hoja = libro.active
    celda = hoja.cell(row=2, column=2)
    valor = str(celda.value)
    formato = celda.number_format
    libro.close()
    assert valor == "600123123"
    assert formato == "@"


